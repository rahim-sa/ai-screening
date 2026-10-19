"""Reads the use case library from the Excel file and filters it."""
from pathlib import Path

from openpyxl import load_workbook

LIBRARY_PATH = Path(__file__).resolve().parents[2] / "data" / "AI_Use_Case_Library.xlsx"
SHEET = "Use Case Master"
CROSS = "Cross-industry"


def load_use_cases(path=LIBRARY_PATH):
    """Return the use cases as a list of dicts (column name -> value)."""
    sheet = load_workbook(path, data_only=True)[SHEET]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    return [row for row in rows if row.get("Specific use case")]


def list_industries(rows):
    """Industries to choose from (without Cross-industry), sorted."""
    return sorted({row["Industry"] for row in rows if row["Industry"] != CROSS})


def use_cases_for(rows, industry):
    """Use cases of this industry plus the cross-industry ones."""
    return [row for row in rows if row["Industry"] in (industry, CROSS)]


def parse_numbers(text, count):
    """'1, 4,7' -> [0, 3, 6]. Returns None if the text is not valid."""
    try:
        numbers = [int(part) for part in text.replace(" ", "").split(",") if part]
    except ValueError:
        return None
    if not numbers or len(set(numbers)) != len(numbers):
        return None
    if any(n < 1 or n > count for n in numbers):
        return None
    return [n - 1 for n in numbers]